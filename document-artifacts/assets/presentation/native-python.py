"""Original editable examples; requires python-pptx in the selected environment."""

import io
import json
import math
from pathlib import Path
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def text(slide, value, box, style, size, color=None):
    frame = slide.shapes.add_textbox(*(Inches(n) for n in box)).text_frame
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    for i, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(0)
        paragraph.font.name = style["font"]
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(color or style["foreground"])


def rectangle(slide, box, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Inches(n) for n in box))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    shape.shadow.inherit = False


def base(slide, style, number):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(style["background"])
    rectangle(slide, (0.7, 0.66, 0.7, 0.07), style["accent"])
    text(slide, f"{number:02}", (12, 6.9, 0.6, 0.3), style, 12, style["muted"])


def title_slide(slide, content, style):
    base(slide, style, 1)
    text(slide, "A COMPOSITION STUDY", (0.7, 1.25, 10, 0.4), style, 14, style["accent"])
    text(slide, content["title"], (0.7, 2, 10.9, 2), style, 48)
    text(slide, content["subtitle"], (0.75, 5.1, 11, 0.7), style, 22, style["muted"])


def comparison_slide(slide, content, style):
    base(slide, style, 2)
    text(slide, content["comparison_title"], (0.7, 1.1, 12, 1.15), style, 32)
    for column, x in ((content["left"], 0.7), (content["right"], 6.9)):
        rectangle(slide, (x, 2.7, 5.7, 3.4), style["panel"])
        text(slide, column["heading"], (x + 0.35, 3.05, 5, 0.7), style, 26, style["accent"])
        text(slide, column["body"], (x + 0.35, 4, 5, 1.65), style, 21)


def closing_slide(slide, content, style):
    base(slide, style, 3)
    text(slide, content["closing_title"], (0.7, 1.6, 11.6, 2.3), style, 44)
    text(slide, content["closing_body"], (0.7, 4.8, 10.5, 1.15), style, 21, style["muted"])


def evidence_slides(deck, content, style):
    """Two bounded native compositions for a small categorical evidence brief.

    Each accepted note gets its own slide; real rendering still decides text fit.
    This is an example content shape, not a generic slide-layout engine.
    """
    evidence = content["evidence"]
    labels, values = evidence["categories"], evidence["values"]
    treatment = content.get("treatment", "briefing")
    if treatment not in ("briefing", "report"):
        raise ValueError("Choose briefing (chart beside table) or report (separate evidence pages)")
    if not 1 <= len(labels) <= 6 or len(labels) != len(values):
        raise ValueError("This example needs one to six paired categories/values; split larger comparisons")
    if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in values):
        raise ValueError("Chart values must be finite numbers")

    def page(heading):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        base(slide, style, len(deck.slides))
        text(slide, heading, (0.7, 1.05, 11.7, 1.5), style, 30)
        return slide

    def chart(slide, box):
        data = CategoryChartData()
        data.categories = labels
        data.add_series(evidence["series"], values)
        obj = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,
                                     *(Inches(n) for n in box), data).chart
        obj.has_legend = False
        obj.font.name, obj.font.size = style["font"], Pt(16)
        obj.font.color.rgb = RGBColor.from_string(style["foreground"])
        obj.category_axis.tick_labels.font.size = Pt(16)
        obj.category_axis.reverse_order = True
        if min(values) >= 0:
            obj.value_axis.minimum_scale = 0
        obj.value_axis.tick_labels.font.size = Pt(13)
        plot = obj.plots[0]
        plot.has_data_labels = True
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.size = Pt(16)
        series = obj.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(style["accent"])
        series.format.line.fill.background()

    def table(slide, box):
        obj = slide.shapes.add_table(len(labels) + 1, 2, *(Inches(n) for n in box)).table
        obj.columns[0].width = Inches(box[2] * .62)
        obj.columns[1].width = Inches(box[2] * .38)
        for row, cells in enumerate([("Category", evidence["unit"]), *zip(labels, map(str, values))]):
            for col, value in enumerate(cells):
                cell = obj.cell(row, col)
                cell.text = value
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(style["panel"])
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name, paragraph.font.size = style["font"], Pt(18)
                    paragraph.font.bold = row == 0
                    paragraph.font.color.rgb = RGBColor.from_string(style["foreground"])

    cover = page(content["title"])
    text(cover, content["subtitle"], (.75, 3.6, 11.5, 1.4), style, 24, style["muted"])
    visual = page(evidence["claim"])
    if treatment == "briefing":
        chart(visual, (.6, 2.55, 7.1, 3.55))
        table(visual, (8.0, 2.65, 4.55, 3.3))
    else:
        chart(visual, (.8, 2.45, 11.7, 3.75))
        detail = page(evidence["claim"])
        table(detail, (1.1, 2.55, 11.1, 3.55))
        text(detail, evidence["source"], (.75, 6.4, 11.3, .4), style, 13, style["muted"])
    text(visual, evidence["source"], (.75, 6.4, 11.3, .4), style, 13, style["muted"])
    for index, note in enumerate(evidence["notes"], 1):
        slide = page(f"{evidence['claim']} / {index}")
        rectangle(slide, (.7, 2.6, .08, 3.35), style["accent"])
        text(slide, note, (1.1, 2.7, 10.9, 3.25), style, 24)


def create_deck(content, style):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = content["title"]
    deck.core_properties.subject = "Synthetic sample content"
    if "evidence" in content:
        evidence_slides(deck, content, style)
    else:
        for layout in (title_slide, comparison_slide, closing_slide):
            layout(deck.slides.add_slide(deck.slide_layouts[6]), content, style)
    return deck


if __name__ == "__main__":
    if len(sys.argv) != 5:
        raise SystemExit("usage: python native-python.py brief.json styles.json style output.pptx")
    brief, styles, name, output = sys.argv[1:]
    content = json.loads(Path(brief).read_text(encoding="utf-8"))
    style = json.loads(Path(styles).read_text(encoding="utf-8"))[name]
    buffer = io.BytesIO()
    create_deck(content, style).save(buffer)
    with open(output, "xb") as stream:
        stream.write(buffer.getvalue())
    print(json.dumps({"output": output, "renderer": "not_invoked"}))
