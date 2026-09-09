"""Replace a unique run, or an explicit --span across runs, in shapes/table cells.

Replacement inherits the first touched run's formatting. Untouched prefixes,
suffixes and runs retain their formatting. This does not edit fields or breaks.
"""

import argparse
import io
import json
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def paragraphs(shapes, shape_id=None):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from paragraphs(shape.shapes, shape_id)
        elif shape_id is None or shape.shape_id == shape_id:
            if shape.has_text_frame:
                yield from shape.text_frame.paragraphs
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if not cell.is_spanned:
                            yield from cell.text_frame.paragraphs


def replace_text(slide, old, new, shape_id=None, span=False):
    if not old or any(char in old + new for char in "\n\r\v"):
        raise ValueError("Use a nonempty paragraph-local span without line breaks")
    matches = []
    for paragraph in paragraphs(slide.shapes, shape_id):
        if not span:
            offset = 0
            for run in paragraph.runs:
                if run.text == old:
                    matches.append((paragraph, offset))
                offset += len(run.text)
            continue
        if old not in paragraph.text:
            continue
        if paragraph.text != "".join(run.text for run in paragraph.runs):
            raise ValueError("Target contains a field or break; use a native-aware edit")
        for match in re.finditer(f"(?={re.escape(old)})", paragraph.text):
            matches.append((paragraph, match.start()))
    if len(matches) != 1:
        raise ValueError("Expected exactly one matching span; inspect or select a shape ID")
    paragraph, start = matches[0]
    end, offset, inserted = start + len(old), 0, False
    for run in paragraph.runs:
        original = run.text
        limit = offset + len(original)
        if offset < end and limit > start:
            prefix = original[:max(0, start - offset)]
            suffix = original[max(0, end - offset):]
            run.text = prefix + (new if not inserted else "") + suffix
            inserted = True
        offset = limit


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("slide_number", type=int)
    parser.add_argument("old_text")
    parser.add_argument("new_text")
    parser.add_argument("output")
    parser.add_argument("--shape-id", type=int, help="Current leaf shape ID, or table shape ID")
    parser.add_argument("--span", action="store_true", help="Match within/across runs in one paragraph")
    args = parser.parse_args()
    deck = Presentation(args.source)
    index = args.slide_number - 1
    if not 0 <= index < len(deck.slides):
        raise SystemExit("Slide number is outside the presentation")
    replace_text(deck.slides[index], args.old_text, args.new_text, args.shape_id, args.span)
    buffer = io.BytesIO()
    deck.save(buffer)
    with open(args.output, "xb") as stream:
        stream.write(buffer.getvalue())
    print(json.dumps({"output": args.output, "slide": index + 1, "renderer": "not_invoked"}))
