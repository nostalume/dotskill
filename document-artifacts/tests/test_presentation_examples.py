"""Native integration: Python with python-pptx and DOTSKILL_TEST_JS_PROJECT for Node.

The explicit JS project must already contain an admitted PptxGenJS installation.
Tests create and remove their own temporary child directory, without installation.
"""

import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

from pptx import Presentation
from pptx.enum.chart import XL_AXIS_CROSSES
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


ASSETS = Path(__file__).resolve().parents[1] / "assets/presentation"
JS_PROJECT = os.environ.get("DOTSKILL_TEST_JS_PROJECT")
NODE = shutil.which("node")


def branded_fixture(output):
    """Original native edit fixture, independent of the maintained slide author."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.7), Inches(.7), Inches(11), Inches(1))
    title.name = "Accepted brand heading"
    paragraph = title.text_frame.paragraphs[0]
    for text, bold in (("Review ", False), ("North", True), (" region", False)):
        run = paragraph.add_run()
        run.text, run.font.bold, run.font.size = text, bold, Pt(32)
    group = slide.shapes.add_group_shape()
    label = group.shapes.add_textbox(Inches(.7), Inches(2), Inches(5), Inches(1))
    label.name = "Grouped annotation"
    label.text = "Keep grouped annotation"
    table = slide.shapes.add_table(2, 2, Inches(.7), Inches(3.4), Inches(8), Inches(1.4)).table
    for row, values in enumerate((("Region", "Count"), ("North", "24"))):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.7), Inches(1), Inches(11), Inches(1)).text = "Independent user note"
    deck.save(output)
    return title.shape_id


def xml_structure(data):
    def node(element):
        return (element.tag, sorted(element.attrib.items()), element.text or "", [node(child) for child in element])
    return node(ET.fromstring(data))


class PresentationFixture(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="native-composition-", dir=self.project)
        self.addCleanup(temporary.cleanup)
        self.area = Path(temporary.name)
        for source in ASSETS.iterdir():
            if source.is_file():
                shutil.copyfile(source, self.area / source.name)

    def author(self, provider, style="paper", output="deck.pptx"):
        command = [NODE, "native-js.mjs"] if provider == "js" else [sys.executable, "-I", "-B", "native-python.py"]
        completed = subprocess.run(
            [*command, "brief.json", "styles.json", style, output],
            cwd=self.area, capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(json.loads(completed.stdout)["renderer"], "not_invoked")
        return self.area / output


class PythonPresentationTests(PresentationFixture):
    project = None

    def test_evidence_treatments_preserve_native_values_and_long_content(self):
        content = json.loads((Path(__file__).parent / "fixtures/evidence.json").read_text(encoding="utf-8"))
        counts = []
        for treatment in ("briefing", "report"):
            content["treatment"] = treatment
            (self.area / "brief.json").write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
            deck = Presentation(self.author("python", output=f"{treatment}.pptx"))
            counts.append(len(deck.slides))
            texts = [s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame]
            for expected in [content["title"], *content["evidence"]["notes"]]:
                self.assertIn(expected, texts)
            charts = [s.chart for slide in deck.slides for s in slide.shapes if s.has_chart]
            tables = [s.table for slide in deck.slides for s in slide.shapes if s.has_table]
            self.assertEqual(len(charts), 1)
            self.assertEqual(list(charts[0].series[0].values), content["evidence"]["values"])
            self.assertEqual([c.label for c in charts[0].plots[0].categories], content["evidence"]["categories"])
            axis = charts[0].value_axis
            self.assertEqual(axis.minimum_scale, 0)
            if axis.crosses == XL_AXIS_CROSSES.CUSTOM:
                self.assertEqual(axis.crosses_at, 0)
            else:
                self.assertIn(axis.crosses, (XL_AXIS_CROSSES.AUTOMATIC, XL_AXIS_CROSSES.MINIMUM))
            self.assertEqual(len(tables), 1)
            self.assertEqual([tables[0].cell(i + 1, 1).text for i in range(3)], ["24", "38", "31"])
        self.assertEqual(counts[1], counts[0] + 1)

    def test_formatted_span_group_and_table_edits_preserve_other_objects(self):
        source = self.area / "brand.pptx"
        title_id = branded_fixture(source)
        original = source.read_bytes()
        changes = [("North region", "Central area", ["--span", "--shape-id", str(title_id)]),
                   ("Keep grouped annotation", "Updated grouped annotation", []),
                   ("24", "25", [])]
        current = source
        for index, (old, new, extra) in enumerate(changes):
            output = self.area / f"revision-{index}.pptx"
            result = subprocess.run(
                [sys.executable, "-I", "-B", "edit-text.py", str(current), "1", old, new, str(output), *extra],
                cwd=self.area, capture_output=True, text=True, timeout=30,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            current = output
        final = Presentation(current)
        title = next(s for s in final.slides[0].shapes if s.shape_id == title_id)
        self.assertEqual(title.text, "Review Central area")
        self.assertFalse(title.text_frame.paragraphs[0].runs[0].font.bold)
        self.assertTrue(title.text_frame.paragraphs[0].runs[1].font.bold)
        group = next(s for s in final.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
        self.assertEqual(group.shapes[0].text, "Updated grouped annotation")
        table = next(s.table for s in final.slides[0].shapes if s.has_table)
        self.assertEqual(table.cell(1, 1).text, "25")
        self.assertEqual(table.cell(1, 0).text, "North")
        self.assertEqual(source.read_bytes(), original)
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(current) as after:
            self.assertEqual(xml_structure(before.read("ppt/slides/slide2.xml")),
                             xml_structure(after.read("ppt/slides/slide2.xml")))

    def test_span_ambiguity_breaks_and_wrong_locator_preserve_source(self):
        source = self.area / "brand.pptx"
        title_id = branded_fixture(source)
        deck = Presentation(source)
        paragraph = deck.slides[0].shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0]
        paragraph.add_run().text = "Before"
        paragraph.add_line_break()
        paragraph.add_run().text = "after"
        deck.save(source)
        original = source.read_bytes()
        for old, options in [("North", ["--span"]),
                             ("Before", ["--span"]),
                             ("North", ["--shape-id", str(title_id + 1000)]),
                             ("", ["--span"])]:
            result = subprocess.run(
                [sys.executable, "-I", "-B", "edit-text.py", str(source), "1", old, "new", "failed.pptx", *options],
                cwd=self.area, capture_output=True, text=True, timeout=30,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse((self.area / "failed.pptx").exists())
            self.assertEqual(source.read_bytes(), original)

    def test_shared_brief_and_styles_produce_native_editable_content(self):
        content = json.loads((self.area / "brief.json").read_text())
        styles = json.loads((self.area / "styles.json").read_text())
        for provider in ("python",):
            for style in styles:
                with self.subTest(provider=provider, style=style):
                    output = self.author(provider, style, f"{provider}-{style}.pptx")
                    deck = Presentation(output)
                    self.assertEqual(len(deck.slides), 3)
                    title = next(s for s in deck.slides[0].shapes if s.has_text_frame and s.text == content["title"])
                    paragraph = title.text_frame.paragraphs[0]
                    color = paragraph.runs[0].font.color
                    if color.type is None:
                        color = paragraph.font.color
                    self.assertEqual(str(color.rgb), styles[style]["foreground"])
                    self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9, places=5)

    def test_output_can_be_edited_without_regenerating_other_slides(self):
        source = self.author("python")
        original = source.read_bytes()
        completed = subprocess.run(
            [sys.executable, "-I", "-B", "edit-text.py", "deck.pptx", "2", "Independent", "Individually useful", "edited.pptx"],
            cwd=self.area, capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(source.read_bytes(), original)
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(self.area / "edited.pptx") as after:
            for part in ("ppt/slides/slide1.xml", "ppt/slides/slide3.xml"):
                self.assertEqual(xml_structure(before.read(part)), xml_structure(after.read(part)))
        edited = Presentation(self.area / "edited.pptx")
        self.assertIn("Individually useful", [s.text for s in edited.slides[1].shapes if s.has_text_frame])

    def test_ambiguous_edit_and_existing_output_do_not_replace_artifacts(self):
        source = self.author("python")
        before = source.read_bytes()
        completed = subprocess.run(
            [sys.executable, "-I", "-B", "edit-text.py", "deck.pptx", "2", "absent", "new", "edited.pptx"],
            cwd=self.area, capture_output=True, text=True, timeout=30,
        )
        self.assertNotEqual(completed.returncode, 0)
        self.assertFalse((self.area / "edited.pptx").exists())
        ambiguous = Presentation(source)
        ambiguous.slides[1].shapes.add_textbox(0, 0, 1000000, 1000000).text = "Independent"
        ambiguous.save(self.area / "ambiguous.pptx")
        completed = subprocess.run(
            [sys.executable, "-I", "-B", "edit-text.py", "ambiguous.pptx", "2", "Independent", "new", "edited.pptx"],
            cwd=self.area, capture_output=True, text=True, timeout=30,
        )
        self.assertNotEqual(completed.returncode, 0)
        self.assertFalse((self.area / "edited.pptx").exists())
        completed = subprocess.run(
            [sys.executable, "-I", "-B", "native-python.py", "brief.json", "styles.json", "night", "deck.pptx"],
            cwd=self.area, capture_output=True, text=True, timeout=30,
        )
        self.assertNotEqual(completed.returncode, 0)
        self.assertEqual(source.read_bytes(), before)


@unittest.skipUnless(JS_PROJECT and NODE, "requires an explicitly admitted JS comparison project")
class JavaScriptPresentationTests(PresentationFixture):
    project = JS_PROJECT

    def test_native_content_and_cross_author_edit(self):
        source = self.author("js")
        original = source.read_bytes()
        deck = Presentation(source)
        content = json.loads((self.area / "brief.json").read_text())
        self.assertEqual(len(deck.slides), 3)
        self.assertIn(content["title"], [s.text for s in deck.slides[0].shapes if s.has_text_frame])
        completed = subprocess.run(
            [sys.executable, "-I", "-B", "edit-text.py", "deck.pptx", "2", "Independent", "Revised", "edited.pptx"],
            cwd=self.area, capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(source.read_bytes(), original)
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(self.area / "edited.pptx") as after:
            for part in ("ppt/slides/slide1.xml", "ppt/slides/slide3.xml"):
                self.assertEqual(xml_structure(before.read(part)), xml_structure(after.read(part)))
        completed = subprocess.run(
            [NODE, "native-js.mjs", "brief.json", "styles.json", "night", "deck.pptx"],
            cwd=self.area, capture_output=True, text=True, timeout=30,
        )
        self.assertNotEqual(completed.returncode, 0)
        self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
